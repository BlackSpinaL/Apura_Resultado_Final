"""
app.py — Boletim Padrão (Streamlit)

Faz upload do arquivo boletim_por_disciplina.xlsx (aba 'Boletins')
e devolve um novo .xlsx com a aba padrão (formato Sugestão),
com todos os alunos.
"""

import io
import streamlit as st
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# =============================================================
# CONFIG
# =============================================================
st.set_page_config(
    page_title="Boletim Padrão",
    page_icon="📘",
    layout="wide",
)

ABA_ENTRADA = "Boletins"
ABA_SAIDA   = "Boletins"   # nome da aba final (padrão)


# =============================================================
# NÚCLEO — TRANSFORMAÇÃO
# =============================================================
def localizar_blocos(ws):
    """Retorna [(linha_cabecalho, linha_resultado), ...] para cada aluno."""
    blocos = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.startswith("MATRÍCULA:"):
            res_row = None
            for rr in range(r, min(r + 60, ws.max_row + 1)):
                for cc in range(1, ws.max_column + 1):
                    if ws.cell(row=rr, column=cc).value == "Resultado":
                        res_row = rr
                        break
                if res_row:
                    break
            if res_row:
                blocos.append((r, res_row))
    return blocos


def escrever_bloco(ws_out, out, ws_in, hrow, res_row):
    d_start = hrow + 3
    d_end   = res_row - 1
    n_disc  = d_end - d_start + 1

    # cabeçalho
    ws_out.cell(row=out, column=1).value  = ws_in.cell(row=hrow, column=1).value
    ws_out.cell(row=out, column=2).value  = ws_in.cell(row=hrow, column=2).value
    ws_out.cell(row=out, column=3).value  = ws_in.cell(row=hrow, column=3).value
    ws_out.cell(row=out, column=12).value = "RELATÓRIO APURA"
    ws_out.cell(row=out, column=13).value = (
        f'=COUNTIF(K{out+3}:K{out+3+n_disc-1},"Reprovado")'
    )
    ws_out.cell(row=out, column=14).value = "Situação - Final de Ano"

    # títulos das etapas
    ws_out.cell(row=out+1, column=1).value  = "Disciplina"
    ws_out.cell(row=out+1, column=2).value  = "1ª Etapa – Média de 18 pontos"
    ws_out.cell(row=out+1, column=4).value  = "2ª Etapa – Média 21 pontos"
    ws_out.cell(row=out+1, column=6).value  = "3ª Etapa – Média 21 pontos"
    ws_out.cell(row=out+1, column=8).value  = "Soma 3 etapas"
    ws_out.cell(row=out+1, column=10).value = "Verificação por Etapa (Mínimo 60% acumulado)"

    # sub-títulos
    ws_out.cell(row=out+2, column=2).value  = "Notas"
    ws_out.cell(row=out+2, column=3).value  = "Situação na 1ª Etapa"
    ws_out.cell(row=out+2, column=4).value  = "Notas"
    ws_out.cell(row=out+2, column=5).value  = "Situação"
    ws_out.cell(row=out+2, column=6).value  = "Notas"
    ws_out.cell(row=out+2, column=7).value  = "Situação"
    ws_out.cell(row=out+2, column=8).value  = "Notas"
    ws_out.cell(row=out+2, column=10).value = "Situação da 1ª Etapa"
    ws_out.cell(row=out+2, column=11).value = "Situação da 2ª Etapa"
    ws_out.cell(row=out+2, column=12).value = "Situação da 3ª Etapa"

    # disciplinas
    for i, sr in enumerate(range(d_start, d_end + 1)):
        dr = out + 3 + i
        ws_out.cell(row=dr, column=1).value  = ws_in.cell(row=sr, column=1).value
        ws_out.cell(row=dr, column=2).value  = ws_in.cell(row=sr, column=2).value
        ws_out.cell(row=dr, column=3).value  = ws_in.cell(row=sr, column=3).value
        ws_out.cell(row=dr, column=4).value  = ws_in.cell(row=sr, column=4).value
        ws_out.cell(row=dr, column=5).value  = ws_in.cell(row=sr, column=5).value
        ws_out.cell(row=dr, column=6).value  = ws_in.cell(row=sr, column=6).value
        ws_out.cell(row=dr, column=7).value  = ws_in.cell(row=sr, column=7).value
        ws_out.cell(row=dr, column=8).value  = f"=B{dr}+D{dr}+F{dr}"
        ws_out.cell(row=dr, column=10).value = "---"
        ws_out.cell(row=dr, column=11).value = f'=IF(B{dr}+D{dr}>=39,"Aprovado","Reprovado")'
        ws_out.cell(row=dr, column=12).value = "---"

    # resultado / motivo
    r_res = out + 3 + n_disc
    r_mot = r_res + 1
    k_start = out + 3
    k_end   = out + 3 + n_disc - 1

    ws_out.cell(row=r_res, column=9).value  = "Resultado"
    ws_out.cell(row=r_res, column=11).value = (
        f'=IF(M{out}>3,"Reprovado",IF(M{out}>0,"Recuperação","Aprovado"))'
    )
    ws_out.cell(row=r_mot, column=9).value  = "Motivo"
    ws_out.cell(row=r_mot, column=11).value = (
        f'=IF(COUNTIF(K{k_start}:K{k_end},"Reprovado")>=4,"Mais que 3 disciplinas",'
        f'IF(COUNTIF(K{k_start}:K{k_end},"Reprovado")=3,"3 disciplinas",'
        f'IF(COUNTIF(K{k_start}:K{k_end},"Reprovado")=2,"2 disciplinas",'
        f'IF(COUNTIF(K{k_start}:K{k_end},"Reprovado")=1,"1 disciplina","Nenhuma disciplina"))))'
    )
    return r_mot + 3


def aplicar_formatacao(ws_out, blocos_saida):
    cor_cabecalho = PatternFill("solid", fgColor="1F4E78")
    cor_etapa     = PatternFill("solid", fgColor="D9E1F2")
    cor_sub       = PatternFill("solid", fgColor="EDEDED")
    cor_result    = PatternFill("solid", fgColor="FFF2CC")

    fonte_branca = Font(color="FFFFFF", bold=True)
    fonte_bold   = Font(bold=True)

    centralizado = Alignment(horizontal="center", vertical="center", wrap_text=True)
    esquerda     = Alignment(horizontal="left",   vertical="center")

    fino  = Side(border_style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)

    for (h, res, mot, n) in blocos_saida:
        d_start = h + 3
        d_end   = h + 3 + n - 1

        for c in (1, 2, 3, 12, 13, 14):
            cel = ws_out.cell(row=h, column=c)
            cel.fill = cor_cabecalho
            cel.font = fonte_branca
            cel.alignment = esquerda

        for c in range(1, 15):
            cel = ws_out.cell(row=h + 1, column=c)
            cel.fill = cor_etapa
            cel.font = fonte_bold
            cel.alignment = centralizado

        for c in range(1, 15):
            cel = ws_out.cell(row=h + 2, column=c)
            cel.fill = cor_sub
            cel.alignment = centralizado

        for r in range(d_start, d_end + 1):
            for c in range(1, 13):
                ws_out.cell(row=r, column=c).border = borda

        for r in (res, mot):
            ws_out.cell(row=r, column=9).font  = fonte_bold
            ws_out.cell(row=r, column=11).fill = cor_result

    larguras = {
        "A": 22, "B": 10, "C": 20, "D": 10, "E": 16, "F": 10, "G": 16,
        "H": 12, "I": 3, "J": 12, "K": 22, "L": 18, "M": 14, "N": 22,
    }
    for col, w in larguras.items():
        ws_out.column_dimensions[col].width = w


def transformar_arquivo(arquivo_bytes):
    """Recebe bytes de um .xlsx e devolve (bytes_saida, num_alunos)."""
    wb_in = load_workbook(io.BytesIO(arquivo_bytes), data_only=False)

    if ABA_ENTRADA not in wb_in.sheetnames:
        raise ValueError(f"A aba '{ABA_ENTRADA}' não foi encontrada no arquivo.")

    ws_in = wb_in[ABA_ENTRADA]

    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = ABA_SAIDA
    ws_out.sheet_view.showGridLines = False

    blocos_in = localizar_blocos(ws_in)
    if not blocos_in:
        raise ValueError("Nenhum aluno encontrado. Verifique o formato do arquivo.")

    out = 1
    blocos_saida = []
    for hrow, res_row in blocos_in:
        n_disc = (res_row - 1) - (hrow + 3) + 1
        proxima = escrever_bloco(ws_out, out, ws_in, hrow, res_row)
        blocos_saida.append((out, out + 3 + n_disc, out + 3 + n_disc + 1, n_disc))
        out = proxima

    aplicar_formatacao(ws_out, blocos_saida)
    ws_out.freeze_panes = "A2"

    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)
    return buf.getvalue(), len(blocos_saida)


# =============================================================
# INTERFACE STREAMLIT
# =============================================================
st.title("📘 Boletim Padrão — Transformador")
st.caption(
    "Envie o arquivo `boletim_por_disciplina.xlsx` (aba **Boletins**) "
    "e receba de volta o mesmo conteúdo no formato **padrão Sugestão**, "
    "com todos os alunos."
)

with st.sidebar:
    st.header("ℹ️ Como usar")
    st.markdown(
        "1. Clique em **Browse files** e envie o `.xlsx` original.\n"
        "2. Aguarde o processamento.\n"
        "3. Clique em **Baixar boletim_padrao.xlsx**.\n\n"
        "O arquivo de saída contém **uma única aba** já no formato padrão, "
        "pronta para impressão / BI."
    )

arquivo = st.file_uploader(
    "Selecione o arquivo Excel (.xlsx)",
    type=["xlsx"],
    accept_multiple_files=False,
)

if arquivo is not None:
    st.info(f"Arquivo recebido: **{arquivo.name}** ({arquivo.size / 1024:.1f} KB)")

    if st.button("🚀 Processar", type="primary"):
        try:
            with st.spinner("Processando alunos..."):
                bytes_saida, n_alunos = transformar_arquivo(arquivo.getvalue())

            st.success(f"✅ {n_alunos} alunos processados com sucesso!")

            nome_saida = "boletim_padrao.xlsx"
            st.download_button(
                label="⬇️ Baixar boletim_padrao.xlsx",
                data=bytes_saida,
                file_name=nome_saida,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )

        except Exception as e:
            st.error(f"❌ Erro ao processar o arquivo: {e}")
            st.exception(e)

st.divider()
st.caption("Feito com Streamlit · openpyxl")